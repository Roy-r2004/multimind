"""Regression tests for the single authoritative "accepted rehab center"
count — the banner's `accepted_facilities_count`, the Phase 2 Eligible
tab's `accepted_only` list/count, and the Excel export must always agree
on exactly which facilities qualify.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.dependencies import AuthContext
from app.db.models import (
    MapsCensusRun,
    MapsCensusStatus,
    MapsClientEligibility,
    MapsPlace,
)
from app.services.scraping.maps_census_service import maps_census_service
from app.services.scraping.maps_export_service import ELIGIBLE_CENTERS_SHEET, maps_export_service


async def _create_run(db, auth: AuthContext) -> MapsCensusRun:
    run = MapsCensusRun(
        organization_id=auth.org_id,
        created_by=auth.user.id,
        country_code="FI",
        country_name="Finland",
        status=MapsCensusStatus.COMPLETED,
    )
    db.add(run)
    await db.flush()
    await db.commit()
    return run


def _accepted_place(run_id: str, google_place_id: str, name: str, **overrides) -> MapsPlace:
    payload = {
        "run_id": run_id,
        "google_place_id": google_place_id,
        "raw_name": name,
        "canonical_name": name,
        "is_relevant": True,
        "client_eligibility": MapsClientEligibility.ELIGIBLE.value,
        "keep_drop_decision": "keep",
        "international_phone_number": "+358 40 000 0000",
    }
    payload.update(overrides)
    return MapsPlace(**payload)


@pytest.mark.asyncio
async def test_banner_count_matches_phase2_list_total(db, auth):
    """Both counters always match: get_run()'s accepted_facilities_count
    and list_places(accepted_only=True)'s meta.total must agree."""
    run = await _create_run(db, auth)
    db.add_all(
        [
            _accepted_place(run.id, "g1", "Eligible One"),
            _accepted_place(run.id, "g2", "Eligible Two", keep_drop_decision="review"),
            # Accepted: phone is optional, never gates acceptance.
            _accepted_place(run.id, "g3", "No Phone", international_phone_number=None),
            # Not accepted: still under review.
            _accepted_place(
                run.id,
                "g4",
                "Still Review",
                client_eligibility=MapsClientEligibility.REVIEW.value,
            ),
        ]
    )
    await db.commit()

    detail = await maps_census_service.get_run(db, auth, run.id)
    listed = await maps_census_service.list_places(db, auth, run.id, accepted_only=True)

    assert detail.accepted_facilities_count == 3
    assert listed.meta.total == 3
    assert detail.accepted_facilities_count == listed.meta.total
    assert {item.google_place_id for item in listed.items} == {"g1", "g2", "g3"}


@pytest.mark.asyncio
async def test_list_runs_count_matches_get_run_count(db, auth):
    """The grouped-query count used by list_runs (avoiding N+1) must agree
    with the single-run count used by get_run for the same run."""
    run = await _create_run(db, auth)
    db.add(_accepted_place(run.id, "g1", "Eligible One"))
    db.add(_accepted_place(run.id, "g2", "Eligible Two"))
    await db.commit()

    runs = await maps_census_service.list_runs(db, auth)
    summary = next(r for r in runs if r.id == run.id)
    detail = await maps_census_service.get_run(db, auth, run.id)

    assert summary.accepted_facilities_count == 2
    assert summary.accepted_facilities_count == detail.accepted_facilities_count


@pytest.mark.asyncio
async def test_pagination_does_not_change_accepted_total(db, auth):
    """Filtering happens in SQL before LIMIT/OFFSET, so the total stays
    accurate regardless of page size/offset."""
    run = await _create_run(db, auth)
    db.add_all(_accepted_place(run.id, f"g{i}", f"Eligible {i}") for i in range(5))
    await db.commit()

    page_one = await maps_census_service.list_places(
        db, auth, run.id, accepted_only=True, limit=2, offset=0
    )
    page_two = await maps_census_service.list_places(
        db, auth, run.id, accepted_only=True, limit=2, offset=2
    )
    page_three = await maps_census_service.list_places(
        db, auth, run.id, accepted_only=True, limit=2, offset=4
    )

    assert page_one.meta.total == 5
    assert page_two.meta.total == 5
    assert page_three.meta.total == 5
    assert len(page_one.items) == 2
    assert len(page_three.items) == 1


@pytest.mark.asyncio
async def test_manual_eligibility_change_updates_count(db, auth):
    """A manual admin override (the same client_eligibility field the
    review endpoint writes) must immediately change the authoritative
    count — no separate counter to go stale."""
    run = await _create_run(db, auth)
    place = _accepted_place(
        run.id,
        "g1",
        "Pending Review",
        client_eligibility=MapsClientEligibility.REVIEW.value,
        keep_drop_decision="review",
    )
    db.add(place)
    await db.commit()

    before = await maps_census_service.get_run(db, auth, run.id)
    assert before.accepted_facilities_count == 0

    # Same action the admin review endpoint's mark_eligible performs.
    place.client_eligibility = MapsClientEligibility.ELIGIBLE.value
    await db.commit()

    after = await maps_census_service.get_run(db, auth, run.id)
    assert after.accepted_facilities_count == 1


@pytest.mark.asyncio
async def test_excel_export_matches_accepted_list_and_count(db, auth):
    """The Excel export, the accepted_only list, and the authoritative
    count must all agree on exactly which facilities qualify."""
    run = await _create_run(db, auth)
    db.add_all(
        [
            _accepted_place(run.id, "g1", "Export Me"),
            _accepted_place(run.id, "g2", "Also Export Me", keep_drop_decision="review"),
            _accepted_place(run.id, "g3", "Not Yet", client_eligibility=MapsClientEligibility.REVIEW.value),
        ]
    )
    await db.commit()

    detail = await maps_census_service.get_run(db, auth, run.id)
    listed = await maps_census_service.list_places(db, auth, run.id, accepted_only=True)
    content, _ = await maps_export_service.build_workbook(db, auth, run.id)
    exported_names = {
        row[0].value
        for row in load_workbook(BytesIO(content))[ELIGIBLE_CENTERS_SHEET].iter_rows(min_row=2)
        if row[0].value
    }

    assert detail.accepted_facilities_count == 2
    assert listed.meta.total == 2
    assert exported_names == {"Export Me", "Also Export Me"}
    assert len(exported_names) == detail.accepted_facilities_count


# --- C-1 end-to-end: a run whose keep/drop finished but whose kept
# candidates have not been enriched yet must report "awaiting_enrichment"
# through the real get_run/list_runs API contract, not "completed" —
# previously this looked identical to "ran and found nothing". ---


@pytest.mark.asyncio
async def test_get_run_reports_awaiting_enrichment_not_completed(db, auth):
    run = await _create_run(db, auth)
    run.processing_state = {"keep_drop_status": "completed"}
    db.add(
        MapsPlace(
            run_id=run.id,
            google_place_id="g1",
            raw_name="Kept Not Yet Enriched",
            canonical_name="Kept Not Yet Enriched",
            is_relevant=True,
            keep_drop_decision="keep",
            enrichment_status="pending",
        )
    )
    await db.commit()

    detail = await maps_census_service.get_run(db, auth, run.id)
    assert detail.phase2_status == "awaiting_enrichment"
    assert detail.accepted_facilities_count == 0  # not enriched, not eligible yet

    runs = await maps_census_service.list_runs(db, auth)
    summary = next(r for r in runs if r.id == run.id)
    assert summary.phase2_status == "awaiting_enrichment"


@pytest.mark.asyncio
async def test_get_run_reports_completed_when_every_candidate_dropped(db, auth):
    """Keep/drop finished and genuinely nothing is left to enrich — this
    must still collapse to "completed", not "awaiting_enrichment"."""
    run = await _create_run(db, auth)
    run.processing_state = {"keep_drop_status": "completed"}
    db.add(
        MapsPlace(
            run_id=run.id,
            google_place_id="g1",
            raw_name="Dropped",
            canonical_name="Dropped",
            is_relevant=False,
            keep_drop_decision="drop",
            enrichment_status="skipped",
        )
    )
    await db.commit()

    detail = await maps_census_service.get_run(db, auth, run.id)
    assert detail.phase2_status == "completed"
    assert detail.accepted_facilities_count == 0


@pytest.mark.asyncio
async def test_get_run_reports_completed_with_accepted_facilities(db, auth):
    run = await _create_run(db, auth)
    run.processing_state = {"keep_drop_status": "completed", "enrichment_status": "completed"}
    db.add(_accepted_place(run.id, "g1", "Fully Enriched And Eligible"))
    await db.commit()

    detail = await maps_census_service.get_run(db, auth, run.id)
    assert detail.phase2_status == "completed"
    assert detail.accepted_facilities_count == 1


@pytest.mark.asyncio
async def test_get_run_reports_pending_for_zero_discovered_candidates(db, auth):
    """A freshly created run with no places at all (e.g. discovery found
    nothing) must report phase2_status pending, not completed/failed."""
    run = await _create_run(db, auth)

    detail = await maps_census_service.get_run(db, auth, run.id)
    assert detail.phase2_status == "pending"
    assert detail.accepted_facilities_count == 0

    runs = await maps_census_service.list_runs(db, auth)
    summary = next(r for r in runs if r.id == run.id)
    assert summary.phase2_status == "pending"
    assert summary.accepted_facilities_count == 0


@pytest.mark.asyncio
async def test_list_runs_awaiting_enrichment_matches_get_run_for_multiple_runs(db, auth):
    """The grouped (N+1-avoiding) query list_runs uses for
    keeps_awaiting_enrichment must agree with get_run's per-run count when
    several runs are mixed in the same listing."""
    awaiting_run = await _create_run(db, auth)
    awaiting_run.processing_state = {"keep_drop_status": "completed"}
    db.add(
        MapsPlace(
            run_id=awaiting_run.id,
            google_place_id="g1",
            raw_name="Kept Not Yet Enriched",
            canonical_name="Kept Not Yet Enriched",
            is_relevant=True,
            keep_drop_decision="keep",
            enrichment_status="pending",
        )
    )
    done_run = await _create_run(db, auth)
    done_run.processing_state = {"keep_drop_status": "completed", "enrichment_status": "completed"}
    db.add(_accepted_place(done_run.id, "g2", "Enriched And Eligible"))
    await db.commit()

    runs = await maps_census_service.list_runs(db, auth)
    by_id = {r.id: r for r in runs}
    assert by_id[awaiting_run.id].phase2_status == "awaiting_enrichment"
    assert by_id[done_run.id].phase2_status == "completed"

    detail_awaiting = await maps_census_service.get_run(db, auth, awaiting_run.id)
    detail_done = await maps_census_service.get_run(db, auth, done_run.id)
    assert detail_awaiting.phase2_status == by_id[awaiting_run.id].phase2_status
    assert detail_done.phase2_status == by_id[done_run.id].phase2_status
