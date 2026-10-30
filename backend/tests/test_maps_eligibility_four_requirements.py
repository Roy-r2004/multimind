"""Regression tests for the updated Phase 2 eligibility requirements.

Mandatory requirements (per the task that removed the old "primary
program" and "phone required" rules):
1. Ownership: private/nonprofit/NGO only, government excluded.
2. Residential/inpatient: an actual overnight program.
3. Addiction treatment: genuine, but NOT required to be the primary
   program.
4. Geographic location: inside the selected country/state.

Phone, website, email, languages, treatment price, and bed count are all
optional and must never independently disqualify an otherwise-eligible
facility.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.dependencies import AuthContext
from app.db.models import (
    MapsCensusRun,
    MapsCensusStatus,
    MapsClientEligibility,
    MapsFacilityType,
    MapsFundingType,
    MapsLifecycleStatus,
    MapsOperatorType,
    MapsOrganizationScope,
    MapsOwnershipStatus,
    MapsPlace,
)
from app.services.scraping.maps_census_service import _phase2_status, maps_census_service
from app.services.scraping.maps_eligibility import compute_client_eligibility
from app.services.scraping.maps_export_service import ELIGIBLE_CENTERS_SHEET, maps_export_service


def make_place(**overrides):
    base = {
        "operating_status": "open",
        "lifecycle_status": MapsLifecycleStatus.DISCOVERED.value,
        "ownership_status": MapsOwnershipStatus.OWNERSHIP_UNKNOWN.value,
        "operator_type": MapsOperatorType.UNKNOWN.value,
        "funding_type": MapsFundingType.UNKNOWN.value,
        "facility_type": MapsFacilityType.UNKNOWN.value,
        "organization_scope": MapsOrganizationScope.UNKNOWN.value,
        "addiction_focus_confirmed": None,
    }
    base.update(overrides)
    from types import SimpleNamespace

    return SimpleNamespace(**base)


def _eligible_fields(**overrides):
    fields = {
        "ownership_status": MapsOwnershipStatus.CONFIRMED_NON_GOVERNMENT.value,
        "organization_scope": MapsOrganizationScope.FACILITY.value,
        "facility_type": MapsFacilityType.RESIDENTIAL_ADDICTION_REHAB.value,
        "addiction_focus_confirmed": True,
    }
    fields.update(overrides)
    return fields


# --- Task 8.1: private residential addiction center with NO phone number
# must still be accepted (phone is optional). ---


def test_residential_center_without_phone_number_is_eligible():
    place = make_place(**_eligible_fields())
    assert compute_client_eligibility(place) == MapsClientEligibility.ELIGIBLE.value
    assert not hasattr(place, "international_phone_number") or not place.international_phone_number


# --- Task 8.3: outpatient-only addiction clinic rejected (unchanged — this
# requirement was never the one removed). ---


def test_outpatient_only_clinic_rejected():
    place = make_place(**_eligible_fields(facility_type=MapsFacilityType.OUTPATIENT_ADDICTION_CENTER.value))
    assert compute_client_eligibility(place) == MapsClientEligibility.EXCLUDED.value


# --- Task 8.4: government-operated inpatient facility rejected (ownership
# requirement untouched). ---


def test_government_operated_inpatient_facility_rejected():
    place = make_place(**_eligible_fields(ownership_status=MapsOwnershipStatus.CONFIRMED_GOVERNMENT.value))
    assert compute_client_eligibility(place) == MapsClientEligibility.EXCLUDED.value


# --- Task 8.5: facility outside the target geography is never silently
# accepted (downgraded to review, not a hard exclude — a verified-polygon
# edge case can still be wrong near a border; a human confirms). ---


def test_facility_outside_target_geography_not_silently_eligible():
    place = make_place(**_eligible_fields(geo_boundary_match="outside_polygon"))
    assert compute_client_eligibility(place) != MapsClientEligibility.ELIGIBLE.value
    assert compute_client_eligibility(place) == MapsClientEligibility.REVIEW.value


# --- Task 8.6: referral-only facility with no own inpatient program is
# rejected. There is no dedicated "referral only" facility_type in the
# enum — such candidates are classified OUTPATIENT_ADDICTION_CENTER (no
# residential program of their own), which is hard-excluded. ---


def test_referral_only_facility_without_own_inpatient_program_rejected():
    place = make_place(
        **_eligible_fields(
            facility_type=MapsFacilityType.OUTPATIENT_ADDICTION_CENTER.value,
            residential_accommodation=False,
        )
    )
    assert compute_client_eligibility(place) == MapsClientEligibility.EXCLUDED.value


# --- Task 8.7: eligible facility missing every optional field (website,
# email, price, languages, bed count) must still be accepted. ---


def test_eligible_facility_missing_all_optional_fields_still_eligible():
    place = make_place(
        **_eligible_fields(),
        official_website=None,
        raw_website=None,
        contact_email=None,
        languages_spoken=None,
        treatment_price=None,
        bed_count=None,
        international_phone_number=None,
    )
    assert compute_client_eligibility(place) == MapsClientEligibility.ELIGIBLE.value


async def _create_run(db, auth: AuthContext) -> MapsCensusRun:
    run = MapsCensusRun(
        organization_id=auth.org_id,
        created_by=auth.user.id,
        country_code="FI",
        country_name="Finland",
        status=MapsCensusStatus.COMPLETED,
    )
    db.add(run)
    await db.flush()
    await db.commit()
    return run


# --- Task 8.1 + 8.8 end-to-end: a phone-less eligible facility reaches the
# accepted API, the Eligible tab list, and the Excel export consistently. ---


@pytest.mark.asyncio
async def test_phoneless_facility_consistent_across_count_list_and_export(db, auth):
    run = await _create_run(db, auth)
    place = MapsPlace(
        run_id=run.id,
        google_place_id="no-phone-1",
        raw_name="Quiet Residential Center",
        canonical_name="Quiet Residential Center",
        is_relevant=True,
        client_eligibility=MapsClientEligibility.ELIGIBLE.value,
        lifecycle_status="confirmed_eligible",
        keep_drop_decision="keep",
        ownership_status=MapsOwnershipStatus.CONFIRMED_NON_GOVERNMENT.value,
        organization_scope=MapsOrganizationScope.FACILITY.value,
        facility_type=MapsFacilityType.RESIDENTIAL_ADDICTION_REHAB.value,
        addiction_focus_confirmed=True,
        international_phone_number=None,
        official_website=None,
    )
    db.add(place)
    await db.commit()

    detail = await maps_census_service.get_run(db, auth, run.id)
    listed = await maps_census_service.list_places(db, auth, run.id, accepted_only=True)
    content, _ = await maps_export_service.build_workbook(db, auth, run.id)
    exported_names = {
        row[0].value
        for row in load_workbook(BytesIO(content))[ELIGIBLE_CENTERS_SHEET].iter_rows(min_row=2)
        if row[0].value
    }

    assert detail.accepted_facilities_count == 1
    assert listed.meta.total == 1
    assert exported_names == {"Quiet Residential Center"}


# --- Task 8.10: Phase 2 empty-state must reflect actual persisted
# execution state, not merely whether the accepted count is zero. ---


def _run_with_state(**overrides):
    base = dict(
        status=MapsCensusStatus.RUNNING,
        processing_state={},
        enrichment_refresh_completed_at=None,
        enrichment_refresh_attempts=0,
        heartbeat_at=None,
    )
    base.update(overrides)
    from types import SimpleNamespace

    return SimpleNamespace(**base)


def test_phase2_status_pending_when_nothing_has_run():
    run = _run_with_state(processing_state={})
    assert _phase2_status(run) == "pending"


def test_phase2_status_running_when_keep_drop_in_progress():
    run = _run_with_state(processing_state={"keep_drop_status": "running"})
    assert _phase2_status(run) == "running"


def test_phase2_status_completed_with_zero_accepted_distinct_from_pending():
    """Phase 2 genuinely completed and found nothing — must report
    "completed", never collapse to the same state as "never ran"."""
    run = _run_with_state(processing_state={"keep_drop_status": "completed"})
    assert _phase2_status(run) == "completed"


def test_phase2_status_completed_when_enrichment_finished():
    run = _run_with_state(
        processing_state={"keep_drop_status": "completed", "enrichment_status": "completed"}
    )
    assert _phase2_status(run) == "completed"


def test_phase2_status_failed_reflects_run_failure():
    run = _run_with_state(status=MapsCensusStatus.FAILED, processing_state={"keep_drop_status": "running"})
    assert _phase2_status(run) == "failed"


# --- F2 follow-up: a keep/drop pass stuck in "running" forever (crashed
# worker, dropped connection) must be distinguishable from one that is
# merely slow, using heartbeat freshness — never by elapsed wall-clock time
# alone, and never by writing to the run just because the UI was opened. ---


def test_phase2_status_running_with_fresh_heartbeat_stays_running():
    from datetime import UTC, datetime

    run = _run_with_state(
        processing_state={
            "keep_drop_status": "running",
            "keep_drop_heartbeat_at": datetime.now(UTC).isoformat(),
        }
    )
    assert _phase2_status(run) == "running"


def test_phase2_status_running_with_stale_heartbeat_is_stalled():
    from datetime import UTC, datetime, timedelta

    stale = datetime.now(UTC) - timedelta(minutes=30)
    run = _run_with_state(
        processing_state={
            "keep_drop_status": "running",
            "keep_drop_heartbeat_at": stale.isoformat(),
        }
    )
    assert _phase2_status(run) == "stalled"


def test_phase2_status_running_with_missing_heartbeat_stays_running():
    """Historical runs recorded before keep_drop_heartbeat_at existed must
    not be guessed at as stalled — absence of heartbeat data is not
    evidence of staleness, only "running" (the prior, unconditional
    behavior) is preserved."""
    run = _run_with_state(processing_state={"keep_drop_status": "running"})
    assert _phase2_status(run) == "running"


def test_phase2_status_does_not_mutate_run_state():
    """Merely computing the status (as happens whenever the UI/API reads a
    run) must never write to processing_state or any other field."""
    from datetime import UTC, datetime, timedelta

    stale = datetime.now(UTC) - timedelta(minutes=30)
    state = {"keep_drop_status": "running", "keep_drop_heartbeat_at": stale.isoformat()}
    run = _run_with_state(processing_state=state)
    before = dict(state)
    assert _phase2_status(run) == "stalled"
    assert run.processing_state == before


# --- C-1 correction: keep/drop completing must never be conflated with
# Phase 2 being fully done when kept candidates are still unenriched (e.g.
# MAPS_CENSUS_AUTO_ENRICHMENT_ENABLED=false and nobody has triggered
# enrichment yet). This previously showed "Phase 2 completed — no
# facilities passed" for a run that had not evaluated anything. ---


def test_phase2_status_awaiting_enrichment_when_keeps_are_unenriched():
    run = _run_with_state(processing_state={"keep_drop_status": "completed"})
    assert _phase2_status(run, keeps_awaiting_enrichment=3) == "awaiting_enrichment"


def test_phase2_status_completed_when_keep_drop_done_and_nothing_awaiting():
    """Every candidate was dropped (or enrichment already finished) — zero
    awaiting is the only case that still collapses to "completed"."""
    run = _run_with_state(processing_state={"keep_drop_status": "completed"})
    assert _phase2_status(run, keeps_awaiting_enrichment=0) == "completed"


def test_phase2_status_completed_overrides_awaiting_count_if_enrichment_done():
    """A stale awaiting count must not fight an explicit enrichment
    completion signal — enrichment_status=="completed" wins."""
    run = _run_with_state(
        processing_state={"keep_drop_status": "completed", "enrichment_status": "completed"}
    )
    assert _phase2_status(run, keeps_awaiting_enrichment=5) == "completed"


def test_phase2_status_stale_failed_enrichment_is_stalled_not_pending():
    """The enrichment watchdog's own stale_failed marker must surface as a
    real stalled/failure signal, not silently collapse to "pending" (as if
    Phase 2 had never started at all)."""
    run = _run_with_state(
        processing_state={"keep_drop_status": "completed", "enrichment_status": "stale_failed"}
    )
    assert _phase2_status(run, keeps_awaiting_enrichment=2) == "stalled"
