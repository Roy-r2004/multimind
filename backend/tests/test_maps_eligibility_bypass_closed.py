"""The strict keep/drop gate decides eligibility, as in the Google Places runs.

A confident "keep" makes a place eligible; Phase 2 classification can only
remove it on hard negative evidence. Delivery to the client export also needs
a phone and a website. These tests exercise the real production functions end
to end: apply_keep_drop -> compute_client_eligibility ->
maps_export_service.build_workbook.
"""

from __future__ import annotations

import json
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.dependencies import AuthContext
from app.db.models import (
    MapsCensusRun,
    MapsCensusStatus,
    MapsClientEligibility,
    MapsFacilityType,
    MapsLifecycleStatus,
    MapsOrganizationScope,
    MapsOwnershipStatus,
    MapsPlace,
)
from app.services.scraping.maps_eligibility import compute_client_eligibility
from app.services.scraping.maps_export_service import (
    ELIGIBLE_CENTERS_SHEET,
    maps_export_service,
)
from app.services.scraping.maps_keep_drop_service import KeepDropDecision, apply_keep_drop

KEEP_JSON = json.dumps(
    {"decision": "keep", "reason": "looks residential", "confidence": 0.92, "evidence": []}
)
REVIEW_JSON = json.dumps(
    {"decision": "review", "reason": "uncertain residential status", "confidence": 0.6, "evidence": []}
)
PHONE = "+358 40 000 0001"
WEBSITE = "https://rehab.example.fi"


def _place(run_id: str, google_place_id: str, name: str, **overrides) -> MapsPlace:
    payload = {
        "run_id": run_id,
        "google_place_id": google_place_id,
        "raw_name": name,
        "canonical_name": name,
    }
    payload.update(overrides)
    return MapsPlace(**payload)


async def _create_run(db, auth: AuthContext) -> MapsCensusRun:
    run = MapsCensusRun(
        organization_id=auth.org_id,
        created_by=auth.user.id,
        country_code="FI",
        country_name="Finland",
        status=MapsCensusStatus.COMPLETED,
    )
    db.add(run)
    await db.flush()
    await db.commit()
    return run


def _completed_classification_fields(**overrides) -> dict:
    fields = {
        "ownership_status": MapsOwnershipStatus.CONFIRMED_NON_GOVERNMENT.value,
        "organization_scope": MapsOrganizationScope.FACILITY.value,
        "facility_type": MapsFacilityType.RESIDENTIAL_ADDICTION_REHAB.value,
        "addiction_focus_confirmed": True,
    }
    fields.update(overrides)
    return fields


async def _exported_names(db, auth, run_id: str) -> list[str]:
    content, _ = await maps_export_service.build_workbook(db, auth, run_id)
    return [
        row[0].value
        for row in load_workbook(BytesIO(content))[ELIGIBLE_CENTERS_SHEET].iter_rows(min_row=2)
        if row[0].value
    ]


@pytest.mark.asyncio
async def test_confident_keep_with_contact_details_exports(db, auth):
    run = await _create_run(db, auth)
    place = _place(run.id, "g1", "Fresh Keep")
    db.add(place)
    await db.commit()

    apply_keep_drop(place, KeepDropDecision.model_validate(json.loads(KEEP_JSON)), source="nano")
    place.international_phone_number = PHONE
    place.raw_website = WEBSITE
    await db.commit()

    assert place.client_eligibility == MapsClientEligibility.ELIGIBLE.value
    assert "Fresh Keep" in await _exported_names(db, auth, run.id)


@pytest.mark.asyncio
async def test_keep_without_website_is_eligible_but_not_exported(db, auth):
    run = await _create_run(db, auth)
    place = _place(run.id, "g2", "Keep Without Website")
    db.add(place)
    await db.commit()

    apply_keep_drop(place, KeepDropDecision.model_validate(json.loads(KEEP_JSON)), source="nano")
    place.international_phone_number = PHONE
    await db.commit()

    assert place.client_eligibility == MapsClientEligibility.ELIGIBLE.value
    assert "Keep Without Website" not in await _exported_names(db, auth, run.id)


@pytest.mark.asyncio
async def test_keep_with_classification_failure_stays_eligible(db, auth):
    """A failed classification is missing evidence, not a disqualification."""
    run = await _create_run(db, auth)
    place = _place(run.id, "g3", "Keep Then Failed Classification")
    db.add(place)
    await db.commit()

    apply_keep_drop(place, KeepDropDecision.model_validate(json.loads(KEEP_JSON)), source="nano")
    place.international_phone_number = PHONE
    place.raw_website = WEBSITE
    place.enrichment_error_message = "classification provider timeout"
    place.client_eligibility = compute_client_eligibility(place)
    await db.commit()

    assert place.client_eligibility == MapsClientEligibility.ELIGIBLE.value
    assert "Keep Then Failed Classification" in await _exported_names(db, auth, run.id)


@pytest.mark.asyncio
async def test_keep_with_completed_eligible_classification_exports(db, auth):
    run = await _create_run(db, auth)
    place = _place(run.id, "g4", "Keep Then Verified Eligible")
    db.add(place)
    await db.commit()

    apply_keep_drop(place, KeepDropDecision.model_validate(json.loads(KEEP_JSON)), source="nano")
    for key, value in _completed_classification_fields().items():
        setattr(place, key, value)
    place.client_eligibility = compute_client_eligibility(place)
    place.international_phone_number = PHONE
    place.raw_website = WEBSITE
    await db.commit()

    assert place.client_eligibility == MapsClientEligibility.ELIGIBLE.value
    assert "Keep Then Verified Eligible" in await _exported_names(db, auth, run.id)


@pytest.mark.asyncio
async def test_keep_with_hard_negative_classification_not_exportable(db, auth):
    run = await _create_run(db, auth)
    place = _place(run.id, "g5", "Keep Then Turned Out Outpatient")
    db.add(place)
    await db.commit()

    apply_keep_drop(place, KeepDropDecision.model_validate(json.loads(KEEP_JSON)), source="nano")
    for key, value in _completed_classification_fields(
        facility_type=MapsFacilityType.OUTPATIENT_ADDICTION_CENTER.value
    ).items():
        setattr(place, key, value)
    place.client_eligibility = compute_client_eligibility(place)
    place.international_phone_number = PHONE
    place.raw_website = WEBSITE
    await db.commit()

    assert place.client_eligibility == MapsClientEligibility.EXCLUDED.value
    assert "Keep Then Turned Out Outpatient" not in await _exported_names(db, auth, run.id)


@pytest.mark.asyncio
async def test_review_with_completed_eligible_classification_exports(db, auth):
    run = await _create_run(db, auth)
    place = _place(run.id, "g6", "Review Then Verified Eligible")
    db.add(place)
    await db.commit()

    apply_keep_drop(
        place, KeepDropDecision.model_validate(json.loads(REVIEW_JSON)), source="sonar"
    )
    assert place.keep_drop_decision == "review"
    assert place.client_eligibility == MapsClientEligibility.REVIEW.value
    for key, value in _completed_classification_fields().items():
        setattr(place, key, value)
    place.client_eligibility = compute_client_eligibility(place)
    place.international_phone_number = PHONE
    place.raw_website = WEBSITE
    await db.commit()

    assert "Review Then Verified Eligible" in await _exported_names(db, auth, run.id)


@pytest.mark.asyncio
async def test_manual_eligible_override_still_exports(db, auth):
    run = await _create_run(db, auth)
    place = _place(
        run.id,
        "g7",
        "Manually Confirmed",
        keep_drop_decision="review",
        lifecycle_status=MapsLifecycleStatus.NEEDS_REVIEW.value,
        client_eligibility=MapsClientEligibility.REVIEW.value,
        international_phone_number=PHONE,
        raw_website=WEBSITE,
    )
    db.add(place)
    await db.commit()

    # Same action the admin review endpoint performs for mark_eligible.
    place.client_eligibility = MapsClientEligibility.ELIGIBLE.value
    place.lifecycle_status = MapsLifecycleStatus.CONFIRMED_ELIGIBLE.value
    await db.commit()

    assert "Manually Confirmed" in await _exported_names(db, auth, run.id)
